"""Langflow component that extracts text and images from a DOCX file.

The text keeps ``<imageN>`` placeholders where images appear in the document,
and the images are returned separately as ``{"imageN": "<base64 PNG>"}``.
A downstream vision model can describe each image and splice the description
back into the text before chunking and indexing.

The file is intentionally self-contained (everything lives inside the component
class): Langflow loads custom components one file at a time, and the code can
also be pasted as-is into the Custom Component editor.
"""

import base64
import os
import re
import shutil
import subprocess
import tempfile
import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from langflow.custom import Component
from langflow.io import BoolInput, DataInput, IntInput, Output
from langflow.schema import Data
from PIL import Image


class DocxParser(Component):
    display_name = 'DOCX Parser'
    description = 'Extract text with <imageN> placeholders and PNG images from a DOCX file.'
    icon = 'file-text'
    name = 'DocxParser'

    inputs = [
        DataInput(
            name='docx_data',
            display_name='DOCX Data',
            info="Data with the file path in 'value' (list of paths) or 'file_path'.",
            tool_mode=True,
        ),
        IntInput(
            name='min_image_size',
            display_name='Min Image Size (px)',
            info='Skip images whose width or height is below this value '
            '(icons, bullets, separators). Set to 0 to keep every image.',
            value=50,
            advanced=True,
        ),
        BoolInput(
            name='include_headers_footers',
            display_name='Include Headers and Footers',
            info='Prepend page headers and append page footers to the text.',
            value=False,
            advanced=True,
        ),
    ]

    outputs = [
        Output(
            name='text_output',
            display_name='Text with image tags',
            method='build_text',
        ),
        Output(
            name='images_output',
            display_name='Images dict',
            method='build_images',
        ),
        Output(
            name='collection_name_output',
            display_name='Collection name',
            method='build_collection_name',
        ),
    ]

    # Raster formats: decoded by Pillow and re-encoded as a clean RGB PNG.
    _RASTER = frozenset(
        {
            'image/png',
            'image/jpeg',
            'image/jpg',
            'image/gif',
            'image/bmp',
            'image/webp',
            'image/tiff',
        }
    )
    # Vector metafiles (typical for diagrams pasted from Visio or old Office):
    # vision models cannot read them, so they are rasterized with LibreOffice.
    _VECTOR = frozenset(
        {
            'image/x-emf',
            'image/emf',
            'image/x-wmf',
            'image/wmf',
        }
    )
    _LIBREOFFICE_TIMEOUT = 120  # seconds per metafile

    # ---------- outputs ----------

    def build_text(self) -> Data:
        text, _ = self._parse_docx()
        path = self._source_path()
        filename = self._display_filename(path) if path else ''
        return Data(data={'filename': filename, 'text': text})

    def build_images(self) -> Data:
        _, images = self._parse_docx()
        return Data(value=images)

    def build_collection_name(self) -> str:
        return str(getattr(self.docx_data, 'collection_name', None) or '')

    # ---------- input handling ----------

    def _source_paths(self) -> list[str]:
        """Input file paths: 'value' (custom loaders) or 'file_path' (File component)."""
        data = self.docx_data
        raw = getattr(data, 'value', None) or getattr(data, 'file_path', None)
        if not raw:
            return []
        if isinstance(raw, (str, os.PathLike)):
            raw = [raw]
        return [os.fspath(p) for p in raw]

    def _source_path(self) -> str | None:
        paths = self._source_paths()
        return paths[0] if paths else None

    @staticmethod
    def _display_filename(path: str) -> str:
        """Last two path components: Langflow stores uploads as '<flow_id>/<file>'."""
        parts = [p for p in re.split(r'[\\/]+', path) if p]
        return '/'.join(parts[-2:])

    def _warn(self, msg: str) -> None:
        try:
            self.log(msg)
        except Exception:
            print(msg)

    # ---------- parsing ----------

    def _parse_docx(self) -> tuple[str, dict[str, str]]:
        paths = self._source_paths()
        if not paths:
            return '', {}

        min_size = max(int(self.min_image_size or 0), 0)
        with_headers = bool(self.include_headers_footers)

        # Every output method calls this on the same instance, and EMF/WMF
        # conversion is expensive, so the result is cached per input.
        key = (paths[0], min_size, with_headers)
        cache = getattr(self, '_parsed_cache', None)
        if cache is not None and cache[0] == key:
            return cache[1], cache[2]

        if len(paths) > 1:
            self._warn(
                f'[DocxParser] {len(paths)} files received, only the first one is parsed: '
                f'{paths[0]}'
            )

        doc = Document(paths[0])
        # images: tag -> base64 PNG; tags: image partname -> tag (None if skipped)
        ctx = SimpleNamespace(images={}, tags={}, min_size=min_size)

        headers: list[str] = []
        footers: list[str] = []
        if with_headers:
            headers, footers = self._render_headers_footers(doc, ctx)
        body = self._render_blocks(doc, ctx)

        # Images that are not referenced inline (floating shapes, unusual
        # anchors) are still returned, numbered after the inline ones.
        for rel in doc.part.rels.values():
            if rel.reltype == RELATIONSHIP_TYPE.IMAGE and not rel.is_external:
                self._tag_for_part(rel.target_part, ctx)

        text = '\n'.join(headers + body + footers)
        self._parsed_cache = (key, text, ctx.images)
        return text, ctx.images

    def _render_blocks(self, container, ctx) -> list[str]:
        """Paragraphs and tables of a document, cell or header, in reading order."""
        lines: list[str] = []
        for block in container.iter_inner_content():
            if isinstance(block, Paragraph):
                lines.append(self._render_paragraph(block, ctx))
            elif isinstance(block, Table):
                lines.extend(self._render_table(block, ctx))
        return lines

    def _render_paragraph(self, paragraph: Paragraph, ctx) -> str:
        chunk: list[str] = []
        # iter_inner_content() also yields hyperlinks, whose runs are not
        # part of paragraph.runs.
        for item in paragraph.iter_inner_content():
            runs = item.runs if isinstance(item, Hyperlink) else [item]
            for run in runs:
                chunk.append(run.text)
                for rid in self._image_rids_in_run(run):
                    tag = self._tag_for_rid(run.part, rid, ctx)
                    if tag is not None:
                        chunk.append(f'<{tag}>')
        return ''.join(chunk)

    def _render_table(self, table: Table, ctx) -> list[str]:
        """Render a table as Markdown, which LLMs read reliably."""
        lines: list[str] = []
        for i, row in enumerate(table.rows):
            cells: list[str] = []
            seen = set()
            for cell in row.cells:
                # A horizontally merged cell is returned once per grid column.
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                text = ' '.join(line for line in self._render_blocks(cell, ctx) if line.strip())
                cells.append(text.replace('\n', ' ').replace('|', '\\|'))
            lines.append('| ' + ' | '.join(cells) + ' |')
            if i == 0:
                lines.append('|' + ' --- |' * len(cells))
        return lines

    def _render_headers_footers(self, doc, ctx) -> tuple[list[str], list[str]]:
        headers: list[str] = []
        footers: list[str] = []
        even_pages = doc.settings.odd_and_even_pages_header_footer
        seen = set()
        for section in doc.sections:
            variants = [('header', headers), ('footer', footers)]
            if section.different_first_page_header_footer:
                variants += [('first_page_header', headers), ('first_page_footer', footers)]
            if even_pages:
                variants += [('even_page_header', headers), ('even_page_footer', footers)]
            for attr, bucket in variants:
                hf = getattr(section, attr)
                # A linked header has no content of its own (checking this
                # first also avoids creating an empty definition).
                if hf.is_linked_to_previous or hf.part.partname in seen:
                    continue
                seen.add(hf.part.partname)
                bucket.extend(line for line in self._render_blocks(hf, ctx) if line.strip())
        return headers, footers

    # ---------- image references ----------

    @staticmethod
    def _image_rids_in_run(run) -> list[str]:
        """rIds of the images in a run: DrawingML (a:blip) and legacy VML (v:imagedata).

        mc:Fallback content is skipped: it duplicates the mc:Choice image for
        old Word versions and would produce a second placeholder.
        """
        not_fallback = "[not(ancestor::*[local-name()='Fallback'])]"
        rids: list[str] = []
        # local-name() avoids depending on the namespace prefixes python-docx registers.
        for blip in run.element.xpath(f".//*[local-name()='blip']{not_fallback}"):
            rid = blip.get(qn('r:embed'))
            if rid and rid not in rids:
                rids.append(rid)
        for imagedata in run.element.xpath(f".//*[local-name()='imagedata']{not_fallback}"):
            rid = imagedata.get(qn('r:id'))
            if rid and rid not in rids:
                rids.append(rid)
        return rids

    def _tag_for_rid(self, part, rid: str, ctx) -> str | None:
        # rIds are scoped to their part: the body and each header have their own.
        image_part = part.related_parts.get(rid)
        if image_part is None:
            return None
        return self._tag_for_part(image_part, ctx)

    def _tag_for_part(self, image_part, ctx) -> str | None:
        """Tag of an image part in reading order; an image used twice shares one tag."""
        key = str(image_part.partname)
        if key not in ctx.tags:
            png = self._normalize_part(image_part, ctx.min_size)
            tag = None
            if png is not None:
                tag = f'image{len(ctx.images) + 1}'
                ctx.images[tag] = png
            ctx.tags[key] = tag
        return ctx.tags[key]

    # ---------- image normalization ----------

    def _normalize_part(self, part, min_size: int = 0) -> str | None:
        """Any DOCX image part -> base64 PNG, or None if unsupported, broken or too small."""
        ct = (part.content_type or '').lower()
        if ct in self._RASTER:
            return self._raster_to_png_b64(part.blob, min_size)
        if ct in self._VECTOR:
            ext = 'wmf' if 'wmf' in ct else 'emf'
            return self._vector_to_png_b64(part.blob, ext, min_size)
        self._warn(f'[DocxParser] unsupported image format skipped: {ct} ({part.partname})')
        return None

    @staticmethod
    def _pil_to_png_b64(im: Image.Image, min_size: int = 0) -> str | None:
        """PIL image -> base64 PNG in RGB, transparency flattened onto white."""
        if im.width < min_size or im.height < min_size:
            return None
        has_alpha = im.mode in ('RGBA', 'LA', 'PA') or (
            im.mode == 'P' and 'transparency' in im.info
        )
        if has_alpha:
            im = im.convert('RGBA')
            bg = Image.new('RGB', im.size, (255, 255, 255))
            bg.paste(im, mask=im.split()[-1])
            im = bg
        elif im.mode != 'RGB':
            im = im.convert('RGB')
        buf = BytesIO()
        im.save(buf, format='PNG')
        return base64.b64encode(buf.getvalue()).decode('ascii')

    def _raster_to_png_b64(self, blob: bytes, min_size: int = 0) -> str | None:
        try:
            with Image.open(BytesIO(blob)) as im:
                im.load()
                return self._pil_to_png_b64(im, min_size)
        except Exception as e:
            self._warn(f'[DocxParser] raster decode failed: {type(e).__name__}: {e}')
            return None

    def _vector_to_png_b64(self, blob: bytes, ext: str, min_size: int = 0) -> str | None:
        soffice = shutil.which('libreoffice') or shutil.which('soffice')
        if not soffice:
            self._warn(
                f'[DocxParser] LibreOffice not found, {ext.upper()} image skipped '
                '(install LibreOffice to extract such diagrams)'
            )
            return None
        with tempfile.TemporaryDirectory() as d:
            src = Path(d, f'img.{ext}')
            src.write_bytes(blob)
            # A separate profile per call: otherwise concurrent conversions
            # block on LibreOffice's single-instance lock.
            profile = Path(d, f'lo_profile_{uuid.uuid4().hex}').as_uri()
            try:
                subprocess.run(
                    [
                        soffice,
                        '--headless',
                        f'-env:UserInstallation={profile}',
                        '--convert-to',
                        'png',
                        '--outdir',
                        d,
                        str(src),
                    ],
                    check=True,
                    capture_output=True,
                    timeout=self._LIBREOFFICE_TIMEOUT,
                )
            except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError) as e:
                self._warn(f'[DocxParser] {ext.upper()} -> PNG conversion failed: {e}')
                return None
            out = Path(d, 'img.png')
            if not out.exists():
                self._warn(f'[DocxParser] {ext.upper()} -> PNG: LibreOffice produced no output')
                return None
            return self._raster_to_png_b64(out.read_bytes(), min_size)
