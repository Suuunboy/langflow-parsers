import re
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches
from langflow.schema import Data

from parsers import docx_parser
from parsers.docx_parser import DocxParser

from conftest import decode_b64, make_png

# ---------- document builders ----------


def add_picture(paragraph, png: bytes) -> None:
    paragraph.add_run().add_picture(BytesIO(png), width=Inches(1))


def add_hyperlink(paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def save(doc, tmp_path, name='doc.docx') -> str:
    path = tmp_path / 'flow-id' / name
    path.parent.mkdir(exist_ok=True)
    doc.save(path)
    return str(path)


def parser_for(path, **kwargs) -> DocxParser:
    data = Data(data={'value': [path], 'collection_name': 'my_collection'})
    return DocxParser(docx_data=data, **kwargs)


def tags_in(text: str) -> list[str]:
    return re.findall(r'<(image\d+)>', text)


# ---------- text and placeholders ----------


def test_text_and_images_in_reading_order(tmp_path):
    doc = Document()
    doc.add_paragraph('Intro')
    p = doc.add_paragraph('Before ')
    add_picture(p, make_png(color=(255, 0, 0)))
    p.add_run(' after')
    add_picture(doc.add_paragraph(), make_png(color=(0, 0, 255)))
    parser = parser_for(save(doc, tmp_path))

    text = parser.build_text().data['text']
    images = parser.build_images().data['value']

    assert text == 'Intro\nBefore <image1> after\n<image2>'
    assert list(images) == ['image1', 'image2']
    assert decode_b64(images['image1']).getpixel((0, 0)) == (255, 0, 0)
    assert decode_b64(images['image2']).getpixel((0, 0)) == (0, 0, 255)


def test_table_rendered_as_markdown_with_images(tmp_path):
    doc = Document()
    doc.add_paragraph('Before table')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Name'
    table.cell(0, 1).text = 'Photo'
    table.cell(1, 0).text = 'a|b'
    add_picture(table.cell(1, 1).paragraphs[0], make_png())
    doc.add_paragraph('After table')
    parser = parser_for(save(doc, tmp_path))

    text = parser.build_text().data['text']

    assert text.splitlines() == [
        'Before table',
        '| Name | Photo |',
        '| --- | --- |',
        '| a\\|b | <image1> |',
        'After table',
    ]
    assert list(parser.build_images().data['value']) == ['image1']


def test_merged_cells_are_not_duplicated(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = 'Merged'
    table.cell(0, 2).text = 'C'
    parser = parser_for(save(doc, tmp_path))

    first_row = parser.build_text().data['text'].splitlines()[0]

    assert first_row == '| Merged | C |'


def test_hyperlink_text_is_kept(tmp_path):
    doc = Document()
    p = doc.add_paragraph('See ')
    add_hyperlink(p, 'the docs', 'https://example.com')
    p.add_run(' for details.')
    parser = parser_for(save(doc, tmp_path))

    assert parser.build_text().data['text'] == 'See the docs for details.'


def test_headers_and_footers_are_optional(tmp_path):
    doc = Document()
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = 'ACME Confidential'
    add_picture(section.header.add_paragraph(), make_png(color=(0, 255, 0)))
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = 'Page footer'
    doc.add_paragraph('Body')
    path = save(doc, tmp_path)

    default = parser_for(path)
    with_hf = parser_for(path, include_headers_footers=True)

    assert default.build_text().data['text'] == 'Body'
    assert default.build_images().data['value'] == {}
    assert with_hf.build_text().data['text'] == 'ACME Confidential\n<image1>\nBody\nPage footer'
    image = decode_b64(with_hf.build_images().data['value']['image1'])
    assert image.getpixel((0, 0)) == (0, 255, 0)


# ---------- image handling ----------


def test_same_image_used_twice_shares_one_tag(tmp_path):
    doc = Document()
    png = make_png()
    add_picture(doc.add_paragraph(), png)
    add_picture(doc.add_paragraph(), png)
    parser = parser_for(save(doc, tmp_path))

    assert tags_in(parser.build_text().data['text']) == ['image1', 'image1']
    assert list(parser.build_images().data['value']) == ['image1']


def test_small_images_are_skipped(tmp_path):
    doc = Document()
    add_picture(doc.add_paragraph('icon '), make_png(10, 10))
    add_picture(doc.add_paragraph('photo '), make_png(200, 100))
    path = save(doc, tmp_path)

    default = parser_for(path)
    keep_all = parser_for(path, min_image_size=0)

    assert default.build_text().data['text'] == 'icon \nphoto <image1>'
    assert decode_b64(default.build_images().data['value']['image1']).size == (200, 100)
    assert tags_in(keep_all.build_text().data['text']) == ['image1', 'image2']


def test_transparency_is_flattened_onto_white(tmp_path):
    doc = Document()
    add_picture(doc.add_paragraph(), make_png(color=(255, 0, 0, 0), mode='RGBA'))
    parser = parser_for(save(doc, tmp_path))

    image = decode_b64(parser.build_images().data['value']['image1'])

    assert image.mode == 'RGB'
    assert image.getpixel((0, 0)) == (255, 255, 255)


def test_unreferenced_and_external_images(tmp_path):
    doc = Document()
    add_picture(doc.add_paragraph(), make_png(color=(255, 0, 0)))
    # An image part with no inline reference (e.g. a floating shape) ...
    doc.part.get_or_add_image(BytesIO(make_png(color=(0, 0, 255))))
    # ... and an externally linked image, which has no data to extract.
    doc.part.relate_to('https://example.com/pic.png', RT.IMAGE, is_external=True)
    parser = parser_for(save(doc, tmp_path))

    images = parser.build_images().data['value']

    assert tags_in(parser.build_text().data['text']) == ['image1']
    assert list(images) == ['image1', 'image2']
    assert decode_b64(images['image2']).getpixel((0, 0)) == (0, 0, 255)


def test_alternate_content_fallback_is_ignored():
    run = parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:v="urn:schemas-microsoft-com:vml"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<mc:AlternateContent>'
        '<mc:Choice Requires="wps"><w:drawing><a:blip r:embed="rId5"/></w:drawing></mc:Choice>'
        '<mc:Fallback><w:pict><v:imagedata r:id="rId6"/></w:pict></mc:Fallback>'
        '</mc:AlternateContent>'
        '</w:r>'
    )

    assert DocxParser._image_rids_in_run(SimpleNamespace(element=run)) == ['rId5']


def test_emf_without_libreoffice_is_skipped_with_warning(monkeypatch, warnings_of):
    monkeypatch.setattr(docx_parser.shutil, 'which', lambda _: None)
    parser = DocxParser()
    messages = warnings_of(parser)
    part = SimpleNamespace(content_type='image/x-emf', blob=b'...', partname='/word/media/a.emf')

    assert parser._normalize_part(part) is None
    assert 'LibreOffice not found' in messages[0]


def fake_libreoffice(monkeypatch, output: bytes | None):
    """Pretend LibreOffice is installed; 'converting' writes `output` as img.png."""
    commands = []

    def run(cmd, **kwargs):
        commands.append(cmd)
        if output is not None:
            outdir = Path(cmd[cmd.index('--outdir') + 1])
            (outdir / 'img.png').write_bytes(output)

    monkeypatch.setattr(docx_parser.shutil, 'which', lambda name: f'/usr/bin/{name}')
    monkeypatch.setattr(docx_parser.subprocess, 'run', run)
    return commands


def test_emf_is_rasterized_with_libreoffice(monkeypatch):
    commands = fake_libreoffice(monkeypatch, make_png(color=(0, 128, 0)))
    part = SimpleNamespace(content_type='image/x-emf', blob=b'...', partname='/word/media/a.emf')

    png = DocxParser()._normalize_part(part)

    assert decode_b64(png).getpixel((0, 0)) == (0, 128, 0)
    cmd = commands[0]
    assert cmd[cmd.index('--convert-to') + 1] == 'png'
    assert cmd[-1].endswith('img.emf')
    # A private LibreOffice profile per call keeps parallel conversions from blocking.
    assert any(arg.startswith('-env:UserInstallation=file:') for arg in cmd)


def test_libreoffice_without_output_is_skipped_with_warning(monkeypatch, warnings_of):
    fake_libreoffice(monkeypatch, output=None)
    parser = DocxParser()
    messages = warnings_of(parser)
    part = SimpleNamespace(content_type='image/x-wmf', blob=b'...', partname='/word/media/a.wmf')

    assert parser._normalize_part(part) is None
    assert 'WMF -> PNG' in messages[0]


def test_unsupported_format_is_skipped_with_warning(warnings_of):
    parser = DocxParser()
    messages = warnings_of(parser)
    part = SimpleNamespace(
        content_type='image/svg+xml', blob=b'<svg/>', partname='/word/media/a.svg'
    )

    assert parser._normalize_part(part) is None
    assert 'unsupported image format' in messages[0]


# ---------- input and outputs ----------


@pytest.mark.parametrize('data', [None, Data(data={'value': []}), Data(data={})])
def test_empty_input(data):
    parser = DocxParser(docx_data=data)

    assert parser.build_text().data == {'filename': '', 'text': ''}
    assert parser.build_images().data['value'] == {}
    assert parser.build_collection_name() == ''


def test_filename_and_collection_name(tmp_path):
    parser = parser_for(save(Document(), tmp_path, 'report.docx'))

    assert parser.build_text().data['filename'] == 'flow-id/report.docx'
    assert parser.build_collection_name() == 'my_collection'


def test_accepts_file_path_from_file_component(tmp_path):
    doc = Document()
    doc.add_paragraph('Hello')
    parser = DocxParser(docx_data=Data(data={'file_path': save(doc, tmp_path)}))

    assert parser.build_text().data['text'] == 'Hello'


def test_only_first_file_is_parsed(tmp_path, warnings_of):
    first, second = Document(), Document()
    first.add_paragraph('first')
    second.add_paragraph('second')
    paths = [save(first, tmp_path, 'a.docx'), save(second, tmp_path, 'b.docx')]
    parser = DocxParser(docx_data=Data(data={'value': paths}))
    messages = warnings_of(parser)

    assert parser.build_text().data['text'] == 'first'
    assert len(messages) == 1
    assert 'only the first one is parsed' in messages[0]


def test_document_is_parsed_once_for_all_outputs(tmp_path, monkeypatch):
    doc = Document()
    doc.add_paragraph('Hello')
    parser = parser_for(save(doc, tmp_path))
    calls = []
    real_document = docx_parser.Document
    monkeypatch.setattr(docx_parser, 'Document', lambda p: calls.append(p) or real_document(p))

    parser.build_text()
    parser.build_images()

    assert len(calls) == 1
